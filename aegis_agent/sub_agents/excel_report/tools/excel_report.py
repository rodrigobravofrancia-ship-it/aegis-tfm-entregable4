import os
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side


# Mapeo de campo JSON -> nombre de columna Excel
COLUMNS = [
    ("Entidad", "entidad_emisora"),
    ("Tipo_norma", "tipo_norma"),
    ("Codigo_norma", "codigo_norma"),
    ("Fecha_publicacion", "fecha_publicacion"),
    ("Glosa", "glosa"),
    ("resumen_norma", "resumen_norma"),
    ("articulos", "articulos"),
    ("area_sugerida", "areas_involucradas"),
    ("justificacion_area", "justificacion_area"),
    ("tipo_alerta", "tipo_alerta"),
    ("justificacion_tipo_alerta", "justificacion_tipo_alerta"),
    ("principales_consideraciones_1", "principales_consideraciones_1"),
    ("principales_consideraciones_2", "principales_consideraciones_2"),
    ("principales_consideraciones_3", "principales_consideraciones_3"),
    ("gravedad", "gravedad"),
    ("justificacion_gravedad", "justificacion_gravedad"),
]


def generate_excel_report(norms: list, zip_name: str, output_dir: str) -> str:
    """Genera un archivo Excel con los campos seleccionados a partir de la lista de normas enriquecidas. Retorna la ruta del archivo generado."""
    total_norms = len(norms)
    print(f"\n--- Generando reporte Excel con {total_norms} normas ---")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Normas"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    headers = [col_name for _, col_name in COLUMNS]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    cell_alignment = Alignment(vertical="top", wrap_text=True)
    for row_idx, norm in enumerate(norms, 2):
        for col_idx, (json_key, _) in enumerate(COLUMNS, 1):
            value = norm.get(json_key, "")
            cell = ws.cell(row=row_idx, column=col_idx, value=str(value))
            cell.alignment = cell_alignment
            cell.border = thin_border

    column_widths = {
        "entidad_emisora": 25,
        "tipo_norma": 20,
        "codigo_norma": 22,
        "fecha_publicacion": 16,
        "glosa": 40,
        "resumen_norma": 45,
        "articulos": 40,
        "areas_involucradas": 25,
        "justificacion_area": 45,
        "tipo_alerta": 16,
        "justificacion_tipo_alerta": 40,
        "principales_consideraciones_1": 40,
        "principales_consideraciones_2": 40,
        "principales_consideraciones_3": 40,
        "gravedad": 14,
        "justificacion_gravedad": 40,
    }
    for col_idx, (_, col_name) in enumerate(COLUMNS, 1):
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = column_widths.get(col_name, 20)

    ws.freeze_panes = "A2"

    os.makedirs(output_dir, exist_ok=True)
    out_path = os.path.join(output_dir, f"reporte_normas_{zip_name}.xlsx")
    wb.save(out_path)
    wb.close()

    print(f"Reporte Excel guardado en: {out_path}\n")
    return out_path
