"""Sub-agente: recibe ZIP + Base_Areas.xlsx (UI o consola), valida y prepara documentos."""
import os
from io import BytesIO
from typing import AsyncGenerator

import openpyxl

from google.adk.agents import BaseAgent
from google.adk.agents.invocation_context import InvocationContext
from google.adk.events import Event, EventActions
from google.genai.types import Content, Part

from .tools.data_loader import run_data_loader


AREAS_EXPECTED_FILENAME = "Base_Areas.xlsx"
REQUIRED_AREAS_HEADERS = ["Grupo", "Subgrupo / Frente", "Temas Asociados"]


class DataLoaderAgent(BaseAgent):
    """PASO 1: Recibe ZIP de normas y Base_Areas.xlsx.

    Soporta dos modos:
      - UI (adk web): los archivos llegan como Parts en ctx.user_content.
      - Consola (main_consola.py): state['zip_path'] viene pre-cargado.

    Persiste ambos archivos en data_staging/, valida estructura, descomprime
    el ZIP y deja en state:
      documents, zip_name, zip_path, areas_excel_path, staging_dir, output_dir.
    """

    async def _run_async_impl(
        self, ctx: InvocationContext
    ) -> AsyncGenerator[Event, None]:
        state = ctx.session.state

        print("=" * 60)
        print("PASO 1: CARGA DE DOCUMENTOS (DataLoader)")
        print("=" * 60)

        project_root = os.path.dirname(
            os.path.dirname(
                os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
            )
        )
        staging_dir = state.get("staging_dir") or os.path.join(project_root, "data_staging")
        output_dir = state.get("output_dir") or os.path.join(project_root, "data_output")
        os.makedirs(staging_dir, exist_ok=True)
        os.makedirs(output_dir, exist_ok=True)

        # ---------- Resolver paths según el modo ----------
        zip_path = state.get("zip_path")
        areas_excel_path = state.get("areas_excel_path")

        if zip_path:
            # Modo consola: zip_path ya está en state. Fallback de areas_excel_path
            # a la misma carpeta donde vive el zip (data/Base_Areas.xlsx por convención).
            if not areas_excel_path:
                areas_excel_path = os.path.join(
                    os.path.dirname(os.path.abspath(zip_path)),
                    AREAS_EXPECTED_FILENAME,
                )
        else:
            # Modo UI: leer ZIP y XLSX desde ctx.user_content
            ok, zip_path, areas_excel_path, err = self._read_uploads_from_user_content(
                ctx, staging_dir
            )
            if not ok:
                yield self._error_event(err)
                return

        # ---------- Validar existencia en disco ----------
        if not os.path.exists(zip_path):
            yield self._error_event(f"El archivo ZIP no existe: {zip_path}")
            return
        if not os.path.exists(areas_excel_path):
            yield self._error_event(
                f"No se encontró Base_Areas.xlsx en: {areas_excel_path}"
            )
            return

        # ---------- Validar estructura del XLSX (ambos modos) ----------
        with open(areas_excel_path, "rb") as f:
            ok, err = self._validate_areas_xlsx(f.read())
        if not ok:
            yield self._error_event(f"Base_Areas.xlsx inválido: {err}")
            return

        # ---------- Descomprimir y listar PDFs ----------
        try:
            loaded = run_data_loader(zip_path)
            documents = loaded.get("documents", [])
        except Exception as e:
            yield self._error_event(f"Error cargando ZIP: {e}")
            return

        zip_name = os.path.splitext(os.path.basename(zip_path))[0]

        msg = f"[DataLoader] {len(documents)} PDFs detectados en {os.path.basename(zip_path)}"
        print(msg)

        yield Event(
            author=self.name,
            content=Content(parts=[Part(text=msg)]),
            actions=EventActions(
                state_delta={
                    "documents": documents,
                    "zip_name": zip_name,
                    "zip_path": zip_path,
                    "staging_dir": staging_dir,
                    "output_dir": output_dir,
                    "areas_excel_path": areas_excel_path,
                    "norms": [],
                    "pipeline_aborted": False,
                }
            ),
        )

    # ----------------- helpers -----------------

    def _read_uploads_from_user_content(
        self, ctx: InvocationContext, staging_dir: str
    ) -> tuple[bool, str, str, str]:
        """Extrae ZIP y Base_Areas.xlsx de ctx.user_content.parts.

        Returns:
            (ok, zip_path, areas_excel_path, error_msg)
        """
        user_content = getattr(ctx, "user_content", None)
        if user_content is None or not getattr(user_content, "parts", None):
            return False, "", "", (
                "No se adjuntaron archivos. Sube el ZIP de normas y el archivo "
                f"{AREAS_EXPECTED_FILENAME} en el chat."
            )

        zip_bytes = None
        zip_filename = None
        xlsx_bytes = None
        xlsx_filename = None

        for part in user_content.parts:
            inline = getattr(part, "inline_data", None)
            if inline is None or not getattr(inline, "data", None):
                continue

            mime = (getattr(inline, "mime_type", "") or "").lower()
            display_name = getattr(inline, "display_name", None) or ""

            is_zip = (
                mime == "application/zip"
                or mime == "application/x-zip-compressed"
                or display_name.lower().endswith(".zip")
            )
            is_xlsx = (
                mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                or display_name.lower().endswith(".xlsx")
            )

            if is_zip:
                if zip_bytes is not None:
                    return False, "", "", "Se adjuntaron varios ZIPs; sube solo uno."
                zip_bytes = inline.data
                zip_filename = display_name or "normas.zip"
            elif is_xlsx:
                if xlsx_bytes is not None:
                    return False, "", "", (
                        f"Se adjuntaron varios .xlsx; sube solo {AREAS_EXPECTED_FILENAME}."
                    )
                xlsx_bytes = inline.data
                xlsx_filename = display_name

        if zip_bytes is None:
            return False, "", "", (
                "Falta el archivo ZIP de normas. Adjunta el ZIP junto con "
                f"{AREAS_EXPECTED_FILENAME}."
            )
        if xlsx_bytes is None:
            return False, "", "", (
                f"Falta el archivo {AREAS_EXPECTED_FILENAME}. Adjúntalo junto con el ZIP."
            )

        # Validar nombre exacto del XLSX (case-sensitive) cuando ADK provee display_name
        if xlsx_filename and xlsx_filename != AREAS_EXPECTED_FILENAME:
            return False, "", "", (
                f"El archivo .xlsx debe llamarse exactamente '{AREAS_EXPECTED_FILENAME}' "
                f"(recibido: '{xlsx_filename}')."
            )

        # Persistir en staging_dir
        zip_path = os.path.join(staging_dir, os.path.basename(zip_filename))
        areas_excel_path = os.path.join(staging_dir, AREAS_EXPECTED_FILENAME)

        with open(zip_path, "wb") as f:
            f.write(zip_bytes)
        with open(areas_excel_path, "wb") as f:
            f.write(xlsx_bytes)

        print(f"[DataLoader] ZIP recibido por UI: {zip_path}")
        print(f"[DataLoader] {AREAS_EXPECTED_FILENAME} recibido por UI: {areas_excel_path}")

        return True, zip_path, areas_excel_path, ""

    @staticmethod
    def _validate_areas_xlsx(data: bytes) -> tuple[bool, str]:
        """Valida estructura mínima de Base_Areas.xlsx."""
        try:
            wb = openpyxl.load_workbook(BytesIO(data), read_only=True)
        except Exception as e:
            return False, f"no se pudo abrir como Excel ({e})"

        try:
            ws = wb.active
            if ws is None:
                return False, "el workbook no tiene hoja activa"
            if ws.max_column < 3:
                return False, f"se esperan al menos 3 columnas, hay {ws.max_column}"
            if ws.max_row < 2:
                return False, "no hay filas de datos (solo header o vacío)"

            header = [c.value for c in ws[1]]
            for i, expected in enumerate(REQUIRED_AREAS_HEADERS):
                actual = header[i] if i < len(header) else None
                if actual != expected:
                    return False, (
                        f"header de columna {i+1} debe ser '{expected}', "
                        f"recibido: '{actual}'"
                    )
        finally:
            wb.close()

        return True, ""

    def _error_event(self, msg: str) -> Event:
        print(f"[DataLoader][ERROR] {msg}")
        return Event(
            author=self.name,
            content=Content(parts=[Part(text=f"ERROR: {msg}")]),
            actions=EventActions(
                state_delta={"pipeline_aborted": True, "pipeline_error": msg}
            ),
        )
